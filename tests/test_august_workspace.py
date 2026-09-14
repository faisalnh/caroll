"""Builder regressions, invoked once by the Node suite or directly via unittest."""
import argparse

import importlib.util
import json
from pathlib import Path
import subprocess
import tempfile
import unittest
import sys
from openpyxl import Workbook
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
SPEC = importlib.util.spec_from_file_location("august_builder", ROOT / "tools/build_august_workspace.py")
B = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(B)
POLICY = {"admin_keys": ["a" * 64], "proposed_admin_key": "b" * 64,
          "regular_exception_key": "c" * 64,
          "reconciliation": {"totals": {"basic": 0, "gross": 0, "deductions": 0},
                             "mismatch_count": 0, "mismatch_metric": "thp",
                             "mismatch_absolute_difference": 1}}
RESULT = "mws-agustus-2026-caroll-results.csv"


class BuilderTests(unittest.TestCase):
    def test_policy_validation_and_evidence(self):
        with tempfile.TemporaryDirectory() as directory:
            config = Path(directory) / "policy.json"
            invalid = [{}, dict(POLICY, admin_keys=[]), dict(POLICY, admin_keys=["invalid"]),
                       dict(POLICY, regular_exception_key="a" * 64), dict(POLICY, extra=True),
                       dict(POLICY, reconciliation={}),
                       dict(POLICY, reconciliation=dict(POLICY["reconciliation"], mismatch_count=True)),
                       dict(POLICY, reconciliation=dict(POLICY["reconciliation"], totals={"basic": 0}))]
            for value in invalid:
                config.write_text(json.dumps(value))
                with self.assertRaises(ValueError):
                    B.load_policy(config)
            config.write_text(json.dumps(POLICY))
            self.assertEqual(B.load_policy(config), POLICY)
        regular, admin = {"1-PM1": 1234567}, {"1-PM1": 1100000}
        for key, code, salary, expected in [
            ("", "1-PM1", 1234567, "regular"), ("", "1-PM1", 1235000, "regular"),
            ("", "1-PM1", 1100000, "admin"), ("", "1-PM2", 1100000, ""),
            ("", "1-PM1", 1200000, ""), ("c" * 64, "1-PM1", 1100000, "regular"),
            ("a" * 64, "1-PM1", 1100000, "admin")]:
            self.assertEqual(B.classify_current(key, code, salary, regular, admin, POLICY)[0], expected)
        self.assertEqual(B.classify_current("", "1-PM1", 1234567, regular, regular, POLICY)[0], "")
        self.assertEqual(B.canonical_golongan("01-pm01"), "1-PM1")
        self.assertEqual(B.money("10.5"), 11)
        self.assertEqual(B.matrix_money("1.2345"), 1235)

    def test_failed_input_removes_stale_artifacts(self):
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            result = folder / RESULT
            result.write_text("stale")
            args = argparse.Namespace(output=folder / "workspace.csv", report=folder / "audit.json",
                                      policy_config=folder / "missing.json")
            args.output.write_text("stale workspace")
            args.report.write_text("stale audit")
            with self.assertRaises(FileNotFoundError):
                B.build(args)
            self.assertTrue(all(not p.exists() for p in (result, args.output, args.report)))
            self.assertFalse(list(folder.glob(".august-build-*")))

    def test_publish_only_after_success(self):
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            args = argparse.Namespace(output=folder / "workspace.csv", report=folder / "audit.json")
            result, staged = folder / RESULT, folder / "staged.csv"
            for returncode, core in [(1, {"totals": {"current": 1}, "issues": []}),
                                     (0, {"totals": None, "issues": []}),
                                     (0, {"totals": {"current": 1}, "issues": [{"severity": "error"}]}),
                                     (0, {"totals": {"current": 1}, "issues": []})]:
                staged.write_text("synthetic result")
                process = subprocess.CompletedProcess([], returncode, json.dumps(core), "")
                with patch.object(B.subprocess, "run", return_value=process):
                    if returncode or not core["totals"] or core["issues"]:
                        with self.assertRaises(ValueError):
                            B.finish_build(args, {}, result, staged)
                        self.assertFalse(result.exists())
                    else:
                        B.finish_build(args, {}, result, staged)
                        self.assertEqual(result.read_text(), "synthetic result")
                        self.assertFalse(staged.exists())

    def check_artifacts(self, folder, config=None):
        completed = subprocess.run(["node", str(ROOT / "tests/august-workspace.test.js"),
                                    "--check-artifacts", str(folder),
                                    str(config or ROOT / "private/august-policy.json")],
                                   cwd=ROOT, capture_output=True, text=True, timeout=60)
        self.assertTrue(completed.returncode == 0, "Private artifact assertions failed; output suppressed")

    def synthetic_inputs(self, folder, unresolved=False):
        matrix = Workbook()
        matrix.remove(matrix.active)
        for name, first, index_row, base in [
            ("Matrix MWS 25_26", 12, 8, 2000),
            ("Draft Matrix MWS 26_27 Penyesua", 13, 9, 2500),
            ("KG 3 Admin", 13, 9, 1500),
        ]:
            sheet = matrix.create_sheet(name)
            for group, column in enumerate((3, 9, 15, 21, 27), 1):
                sheet.cell(7, column, base * group)
                sheet.cell(6, column + 3, 0)
                sheet.cell(index_row, column, 0)
                for kmk in range(15):
                    for tier in range(4):
                        sheet.cell(first + kmk, column + tier, base * group)
        matrix.save(folder / "matrix.xlsx")
        payroll = Workbook()
        sheet = payroll.active
        sheet.title = "0826 MWS"
        bpjs = payroll.create_sheet("BPJSTK")
        people = [("Synthetic Regular", "1-PM1", 2000000),
                  ("Synthetic Admin", "3-PM1", 4000000)]
        if unresolved:
            people.append(("Synthetic Unresolved", "2-PM1", 3333333))
        basic = 0
        for row, (name, code, salary) in enumerate(people, 7):
            employee_id = "SYN-" + str(row)
            for col, value in {"C": employee_id, "D": name, "H": "Synthetic",
                               "I": "Staff", "M": code, "N": salary, "AE": "TK/0",
                               "P": 100000, "AX": 10000, "AN": salary + 100000,
                               "BJ": 10000, "BL": 0, "BM": salary + 90000}.items():
                sheet[f"{col}{row}"] = value
            bpjs.cell(row, 2, employee_id)
            bpjs.cell(row, 12, 20000)
            basic += salary
        payroll.save(folder / "payroll.xlsx")
        kmk = Workbook()
        kmk.active.title = " MWS No Covid-izar"
        kmk.active.cell(2, 1, "Synthetic Regular")
        kmk.active.cell(2, 14, "1-PM2")
        kmk.save(folder / "kmk.xlsx")
        policy = dict(POLICY, admin_keys=[B.policy_key("Synthetic Admin")],
                      reconciliation=dict(POLICY["reconciliation"], totals={
                          "basic": basic, "gross": basic + len(people) * 100000,
                          "deductions": len(people) * 10000}))
        config = folder / "policy.json"
        config.write_text(json.dumps(policy))
        return config

    def run_synthetic_cli(self, folder, config):
        command = [sys.executable, "-B", str(ROOT / "tools/build_august_workspace.py"),
                   "--policy-config", str(config),
                   "--output", str(folder / "mws-agustus-2026-caroll-workspace.csv"),
                   "--report", str(folder / "mws-agustus-2026-audit.json")]
        for source in ("payroll", "kmk", "matrix"):
            command.extend(["--" + source, str(folder / (source + ".xlsx"))])
        return subprocess.run(command, cwd=folder, capture_output=True, text=True, timeout=120)

    def test_reconciliation_failure_publishes_fresh_draft_only(self):
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            config = self.synthetic_inputs(folder)
            policy = json.loads(config.read_text())
            policy["reconciliation"]["totals"]["basic"] += 1
            config.write_text(json.dumps(policy))
            (folder / RESULT).write_text("stale")
            completed = self.run_synthetic_cli(folder, config)
            self.assertEqual(completed.returncode, 1)
            audit = json.loads((folder / "mws-agustus-2026-audit.json").read_text())
            self.assertFalse(audit["reconciliation_passed"])
            self.assertFalse(audit["results_exported"])
            self.assertTrue((folder / "mws-agustus-2026-caroll-workspace.csv").is_file())
            self.assertFalse((folder / RESULT).exists())

    def test_publication_failure_removes_partial_artifacts(self):
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            config = self.synthetic_inputs(folder)
            args = argparse.Namespace(output=folder / "workspace.csv", report=folder / "audit.json",
                                      policy_config=config, payroll=folder / "payroll.xlsx",
                                      kmk=folder / "kmk.xlsx", matrix=folder / "matrix.xlsx")
            replace = B.os.replace

            def fail_audit(source, destination):
                if destination == args.report:
                    raise OSError("Synthetic publication failure")
                return replace(source, destination)

            with patch.object(B.os, "replace", side_effect=fail_audit):
                with self.assertRaises(OSError):
                    B.build(args)
            self.assertTrue(all(not p.exists() for p in (args.output, args.report, folder / RESULT)))
            self.assertFalse(list(folder.glob(".august-*")))

    def test_synthetic_cli_end_to_end(self):
        for unresolved in (False, True):
            with self.subTest(unresolved=unresolved), tempfile.TemporaryDirectory() as directory:
                folder = Path(directory)
                config = self.synthetic_inputs(folder, unresolved)
                for name in (RESULT, "mws-agustus-2026-caroll-workspace.csv", "mws-agustus-2026-audit.json"):
                    (folder / name).write_text("stale")
                completed = self.run_synthetic_cli(folder, config)
                self.assertEqual(completed.returncode, int(unresolved), completed.stderr)
                self.check_artifacts(folder, config)
                audit = json.loads((folder / "mws-agustus-2026-audit.json").read_text())
                self.assertEqual(audit["unresolved_count"], int(unresolved))
                self.assertEqual(audit["calculated_golongan_matches"], 1)
                self.assertEqual(audit["workspace_csv"], str(folder / "mws-agustus-2026-caroll-workspace.csv"))
                self.assertFalse(list(folder.glob(".august-*")))
                # The same destinations must not survive a subsequent early failure.
                (folder / "matrix.xlsx").write_text("invalid XLSX")
                failed = self.run_synthetic_cli(folder, config)
                self.assertEqual(failed.returncode, 1)
                for name in (RESULT, "mws-agustus-2026-caroll-workspace.csv", "mws-agustus-2026-audit.json"):
                    self.assertFalse((folder / name).exists())

    def test_existing_artifacts_without_sources(self):
        folder = ROOT / "private"
        if not (folder / "august-policy.json").is_file():
            self.skipTest("Private config absent")
        if not all((folder / name).exists() for name in ["mws-agustus-2026-caroll-workspace.csv", "mws-agustus-2026-audit.json"]):
            self.skipTest("Private artifacts absent")
        self.check_artifacts(folder)

    def test_current_builder_with_private_sources(self):
        config = ROOT / "private/august-policy.json"
        if not config.exists():
            self.skipTest("Private config absent")
        policy = B.load_policy(config)
        sources = policy.get("source_paths", {})
        if not sources or not all(Path(value).is_file() for value in sources.values()):
            self.skipTest("Authoritative private sources absent")
        with tempfile.TemporaryDirectory(prefix="august-test-", dir=ROOT / "private") as directory:
            folder = Path(directory)
            command = ["python3", "-B", str(ROOT / "tools/build_august_workspace.py"),
                       "--policy-config", str(config), "--output", str(folder / "mws-agustus-2026-caroll-workspace.csv"),
                       "--report", str(folder / "mws-agustus-2026-audit.json")]
            for key, value in sources.items():
                command.extend(["--" + key, value])
            (folder / RESULT).write_text("stale synthetic result")
            completed = subprocess.run(command, cwd=ROOT, capture_output=True, text=True, timeout=120)
            audit_file = folder / "mws-agustus-2026-audit.json"
            self.assertTrue(audit_file.exists(), "Current builder failed before writing audit; output suppressed")
            audit = json.loads(audit_file.read_text())
            self.assertTrue(completed.returncode == (0 if audit.get("results_exported") else 1))
            self.check_artifacts(folder)


if __name__ == "__main__":
    unittest.main()
