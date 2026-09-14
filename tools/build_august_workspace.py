#!/usr/bin/env python3
"""Build a private Caroll workspace from the August 2026 MWS payroll files.

Authoritative local invocation (openpyxl already installed; no new dependencies):
python3 -B tools/build_august_workspace.py \
  --payroll /path/to/payroll.xlsx \
  --kmk /path/to/kmk.xlsx \
  --matrix /path/to/matrix.xlsx \
  --policy-config private/august-policy.json \
  --output private/mws-agustus-2026-caroll-workspace.csv --report private/mws-agustus-2026-audit.json

Only build status goes to stdout. The audit contains private evidence.
Unresolved types remain blank: draft import is allowed, payroll export is not.
"""

from __future__ import annotations

import argparse
import csv
import json
import re
import hashlib
import subprocess
import tempfile
import os

from collections import Counter
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string


WORKSPACE_HEADERS = [
    "schema_version", "record_type", "record_id", "name", "current_period",
    "proposed_period", "matrix_id", "scenario", "effective_date", "generator_settings", "salary_group",
    "professional_category", "kmk_level", "golongan", "basic_salary", "note",
    "employee_id", "current_golongan", "unit", "department", "position",
    "employment_status", "join_date", "active", "proposed_golongan",
    "current_basic_override", "proposed_basic_override", "ptkp_status",
    "bpjs_kesehatan", "bpjs_ketenagakerjaan", "pph_method", "pph_rate",
    "pph_fixed_override", "notes", "code", "category", "direction",
    "calculation_type", "default_value", "taxable", "applies_current",
    "applies_proposed", "rounding", "component_code", "current_value",
    "proposed_value", "employee_rate", "employer_rate", "minimum_basis",
    "maximum_basis", "basis", "employee_enabled", "employer_enabled", "currency",
    "percentage_precision", "include_inactive", "allow_negative_thp",
    "proposed_defaults_current", "tax_enabled", "tax_basis", "tax_rounding",
    "matrix_type", "current_matrix_type", "proposed_matrix_type",
    "proposed_matrix_type_defaults_current",
]

COMPONENTS = [
    ("TUNJ_STRUKTURAL", "Tunjangan struktural", "earning", "P"),
    ("TUNJ_FUNGSIONAL", "Tunjangan fungsional", "earning", "U"),
    ("TRANSPORT", "Transport", "earning", "V"),
    ("KETEPATAN_WAKTU", "Ketepatan waktu", "earning", "W"),
    ("TUNJ_PASANGAN", "Tunjangan pasangan", "earning", "Z"),
    ("TUNJ_ANAK", "Tunjangan anak", "earning", "AC"),
    ("KOREKSI", "Koreksi payroll", "earning", "AF"),
    ("TUNJ_ASURANSI", "Tunjangan asuransi", "earning", "AG"),
    ("TUNJ_TELEPON", "Tunjangan telepon", "earning", "AH"),
    ("TUNJ_BENSIN", "Tunjangan bensin", "earning", "AI"),
    ("SERVICE_KENDARAAN", "Service kendaraan", "earning", "AJ"),
    ("WFH", "Tunjangan WFH", "earning", "AK"),
    ("LEMBUR", "Lembur", "earning", "AM"),
    ("INHEALTH", "Potongan InHealth", "employee_deduction", "AX"),
    ("BPJS_TK_KARYAWAN", "BPJS TK karyawan (aktual Agustus)", "employee_deduction", "AY"),
    ("TOKO_LAZMART", "Toko LazMart", "employee_deduction", "BA"),
    ("U_SEK", "U-Sek", "employee_deduction", "BB"),
    ("CASHBON", "Cashbon", "employee_deduction", "BC"),
    ("POTONGAN_LAIN", "Potongan lain-lain", "employee_deduction", "BD"),
    ("SHARING_CARE", "Sharing & Care", "employee_deduction", "BF"),
    ("SEDEKAH", "Sedekah", "employee_deduction", "BG"),
    ("ZAKAT_SEDEKAH", "Zakat / sedekah tambahan", "employee_deduction", "BL"),
]


def money(value: object) -> int:
    if value in (None, ""):
        return 0
    if isinstance(value, bool):
        raise ValueError("Boolean is not a monetary value")
    return int(Decimal(str(value)).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def matrix_money(value: object, rounding: int = 1) -> int:
    amount = Decimal(str(value)) * Decimal("1000")
    step = Decimal(str(rounding))
    return int((amount / step).quantize(Decimal("1"), rounding=ROUND_HALF_UP) * step)


def normalize_name(value: object) -> str:
    base = str(value or "").lower().split(",", 1)[0]
    return re.sub(r"[^a-z0-9]", "", base)


def parse_golongan(code: str) -> tuple[str, str, int]:
    match = re.fullmatch(r"(\d+)-([A-Za-z]+)(\d+)", str(code).strip())
    if not match:
        raise ValueError(f"Invalid Golongan: {code}")
    return str(int(match.group(1))), match.group(2).upper(), int(match.group(3))


def empty_row(record_type: str, record_id: str) -> dict[str, object]:
    row = {header: "" for header in WORKSPACE_HEADERS}
    row.update(schema_version=1, record_type=record_type, record_id=record_id)
    return row


def source_matrix_lookup(workbook, sheet_name: str, first_data_row: int) -> dict[str, int]:
    sheet = workbook[sheet_name]
    lookup: dict[str, int] = {}
    for group, start_column in enumerate((3, 9, 15, 21, 27), start=1):
        for kmk in range(1, 16):
            for category, offset in (("PM", 0), ("P", 1), ("M", 2), ("U", 3)):
                value = sheet.cell(first_data_row + kmk - 1, start_column + offset).value
                if isinstance(value, (int, float, Decimal)):
                    lookup[f"{group}-{category}{kmk}"] = matrix_money(value)
    return lookup


def decimal_text(value: Decimal) -> str:
    text = format(value, "f").rstrip("0").rstrip(".")
    return text or "0"


def matrix_settings(workbook, sheet_name: str, kmk_row: int) -> list[dict[str, object]]:
    sheet = workbook[sheet_name]
    settings = []
    for start_column in (3, 9, 15, 21, 27):
        base = Decimal(str(sheet.cell(7, start_column).value)) * Decimal("1000")
        cola = Decimal(str(sheet.cell(6, start_column + 3).value))
        kmk_index = Decimal(str(sheet.cell(kmk_row, start_column).value))
        settings.append({
            "base_salary": int(base) if base == base.to_integral_value() else decimal_text(base),
            "cola": decimal_text(cola),
            "kmk_index": decimal_text(kmk_index),
        })
    return settings


def generated_matrix(settings: list[dict[str, object]], rounding: int = 1) -> dict[str, int]:
    lookup: dict[str, int] = {}
    step = Decimal(str(rounding))
    for group, setting in enumerate(settings, start=1):
        base = Decimal(str(setting["base_salary"]))
        cola = Decimal(str(setting["cola"]))
        kmk_index = Decimal(str(setting["kmk_index"]))
        for tier, category in enumerate(("PM", "P", "M", "U")):
            for kmk in range(1, 16):
                exponent = kmk - 1 + 2 * tier
                salary = base * (Decimal("1") + cola) * (Decimal("1") + kmk_index) ** exponent
                rounded = (salary / step).quantize(Decimal("1"), rounding=ROUND_HALF_UP) * step
                lookup[f"{group}-{category}{kmk}"] = int(rounded)
    return lookup


def calculated_golongan(source_path: Path) -> dict[str, str]:
    sheet = load_workbook(source_path, read_only=True, data_only=True)[" MWS No Covid-izar"]
    lookup: dict[str, str] = {}
    for row in sheet.iter_rows(min_row=2, max_row=840, min_col=1, max_col=14, values_only=True):
        name, calculated = row[0], row[13]
        if name and calculated:
            lookup[normalize_name(name)] = str(calculated).strip()
    return lookup


def bpjs_employer_lookup(payroll_book) -> dict[str, int]:
    """Read the employer BPJS amount from the same BPJSTK sheet used by payroll."""
    sheet = payroll_book["BPJSTK"]
    lookup: dict[str, int] = {}
    for row in sheet.iter_rows(min_row=4, max_row=172, min_col=1, max_col=13, values_only=True):
        employee_id = row[1]
        employer_amount = row[11]
        if employee_id and isinstance(employer_amount, (int, float)):
            lookup[str(employee_id).strip()] = money(employer_amount)
    return lookup


def load_policy(path: Path) -> dict:
    """Require explicit private policy; hashes are identifiers, not anonymization."""
    policy = json.loads(path.read_text(encoding="utf-8"))
    required = {"admin_keys", "proposed_admin_key", "regular_exception_key", "reconciliation"}
    if not isinstance(policy, dict) or not required <= policy.keys() or policy.keys() - required - {"source_paths"}:
        raise ValueError("Invalid private policy fields")
    admins = policy["admin_keys"]
    if not isinstance(admins, list) or not admins:
        raise ValueError("Explicit admin policy required")
    keys = admins + [policy["proposed_admin_key"], policy["regular_exception_key"]]
    if any(not isinstance(key, str) or not re.fullmatch(r"[0-9a-f]{64}", key) for key in keys):
        raise ValueError("Invalid private policy identifiers")
    if len(set(keys)) != len(keys):
        raise ValueError("Duplicate or conflicting private policy identifiers")
    if "source_paths" in policy and (
        not isinstance(policy["source_paths"], dict)
        or set(policy["source_paths"]) != {"payroll", "kmk", "matrix"}
        or any(not isinstance(value, str) or not value.strip() for value in policy["source_paths"].values())
    ):
        raise ValueError("Invalid private source paths")
    profile = policy["reconciliation"]
    if (not isinstance(profile, dict)
        or set(profile) != {"totals", "mismatch_count", "mismatch_metric", "mismatch_absolute_difference"}
        or not isinstance(profile["totals"], dict)
        or set(profile["totals"]) != {"basic", "gross", "deductions"}
        or any(type(value) is not int or value < 0 for value in profile["totals"].values())
        or type(profile["mismatch_count"]) is not int or profile["mismatch_count"] < 0
        or profile["mismatch_metric"] != "thp"
        or type(profile["mismatch_absolute_difference"]) is not int
        or profile["mismatch_absolute_difference"] < 1):
        raise ValueError("Invalid private reconciliation profile")
    return policy


def reconciliation_passed(totals, mismatches, profile):
    return (all(totals[key]["source"] == totals[key]["mapped"] == expected
                for key, expected in profile["totals"].items())
            and len(mismatches) == profile["mismatch_count"]
            and all(m["metric"] == profile["mismatch_metric"]
                    and abs(m["difference"]) == profile["mismatch_absolute_difference"]
                    for m in mismatches))


def policy_key(name: object) -> str:
    return hashlib.sha256(normalize_name(name).encode("utf-8")).hexdigest()


def canonical_golongan(code: object) -> str:
    return "{}-{}{}".format(*parse_golongan(str(code)))


def valid_payroll_row(values) -> bool:
    return bool(len(values) > 13 and values[2] and values[3] and values[12]
                and isinstance(values[13], (int, float)) and not isinstance(values[13], bool))


def classify_current(key, code, actual, regular, admin, policy):
    if key == policy["regular_exception_key"]:
        return "regular", "Explicit user regular exception; payroll override until policy verified"
    if key in policy["admin_keys"]:
        return "admin", "Explicit user-confirmed admin; current cell evidenced by August payroll"
    regular_salary = regular.get(code)
    regular_match = regular_salary is not None and (
        regular_salary == actual or money(Decimal(regular_salary) / 1000) * 1000 == actual)
    admin_match = admin.get(code) == actual
    if regular_match and not admin_match:
        return "regular", "Actual basic matches same-Golongan regular cell (exact or payroll nearest-thousand rounding)"
    if admin_match and not regular_match:
        return "admin", "Actual basic matches same-Golongan confirmed admin payroll cell only"
    return "", "Unresolved: ambiguous or unmatched actual basic; job title is not evidence"


def build(args: argparse.Namespace) -> dict[str, object]:
    results_path = args.output.parent / "mws-agustus-2026-caroll-results.csv"
    destinations = (args.output, args.report, results_path)
    if len({path.resolve() for path in destinations}) != len(destinations):
        raise ValueError("Artifact paths must be distinct")
    sources = [getattr(args, key, None) for key in ("policy_config", "payroll", "kmk", "matrix")]
    if any(source and source.resolve() in {p.resolve() for p in destinations} for source in sources):
        raise ValueError("Artifacts must not overwrite inputs")
    for destination in destinations:
        destination.parent.mkdir(parents=True, exist_ok=True)
        destination.unlink(missing_ok=True)
    with tempfile.TemporaryDirectory(prefix=".august-build-", dir=args.output.parent) as directory:
        folder = Path(directory)
        staged = argparse.Namespace(**vars(args))
        staged.output, staged.report = folder / "workspace.csv", folder / "audit.json"
        staged_results = folder / results_path.name
        try:
            report = build_staged(staged)
        finally:
            # Publish only a fresh, complete draft/audit pair, including blocked builds.
            if staged.report.is_file():
                report = json.loads(staged.report.read_text(encoding="utf-8"))
                report.update(workspace_csv=str(args.output), audit_json=str(args.report),
                              results_csv=str(results_path))
                staged.report.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
                try:
                    os.replace(staged.output, args.output)
                    if report["results_exported"]:
                        os.replace(staged_results, results_path)
                    os.replace(staged.report, args.report)
                except Exception:
                    for destination in destinations:
                        destination.unlink(missing_ok=True)
                    raise
        return report


def build_staged(args: argparse.Namespace) -> dict[str, object]:
    results_path = args.output.parent / "mws-agustus-2026-caroll-results.csv"
    policy = load_policy(args.policy_config)
    payroll_book = load_workbook(args.payroll, read_only=True, data_only=True)
    payroll_sheet = payroll_book["0826 MWS"]
    matrix_book = load_workbook(args.matrix, data_only=True)
    current_settings = matrix_settings(matrix_book, "Matrix MWS 25_26", 8)
    proposed_settings = matrix_settings(matrix_book, "Draft Matrix MWS 26_27 Penyesua", 9)
    admin_settings = matrix_settings(matrix_book, "KG 3 Admin", 9)
    current_matrix = generated_matrix(current_settings)
    proposed_matrix = generated_matrix(proposed_settings)
    # KG 3 is the only admin group supported by this source sheet. Its cached
    # cells are authoritative even where the generic generator differs.
    admin_matrix = {code: salary for code, salary in
                    source_matrix_lookup(matrix_book, "KG 3 Admin", 13).items()
                    if code.startswith("3-")}
    if len(admin_matrix) != 60:
        raise ValueError("Incomplete proposed admin source matrix")
    current_admin_matrix = {}
    admin_generator = generated_matrix(admin_settings)
    admin_evidence = {}
    payroll_rows = list(payroll_sheet.iter_rows(min_row=7, max_row=1000, values_only=True))
    for excel_row, values in enumerate(payroll_rows, 7):
        if valid_payroll_row(values) and policy_key(values[3]) in policy["admin_keys"]:
            code = canonical_golongan(values[12])
            salary = money(values[13])
            if code in current_admin_matrix and current_admin_matrix[code] != salary:
                raise ValueError("Conflicting confirmed current admin payroll cells")
            current_admin_matrix[code] = salary
            admin_evidence.setdefault(code, []).append({"row": excel_row, "cell": f"N{excel_row}",
                                                       "employee_id": str(values[2]).strip()})
    source_current_matrix = source_matrix_lookup(matrix_book, "Matrix MWS 25_26", 12)
    source_proposed_matrix = source_matrix_lookup(matrix_book, "Draft Matrix MWS 26_27 Penyesua", 13)
    matrix_mismatches = {
        "current": [code for code, salary in current_matrix.items() if source_current_matrix.get(code) != salary],
        "proposed": [code for code, salary in proposed_matrix.items() if source_proposed_matrix.get(code) != salary],
    }
    if matrix_mismatches["current"] or matrix_mismatches["proposed"]:
        raise ValueError(f"Generated matrix differs from workbook cells: {matrix_mismatches}")
    calculated = calculated_golongan(args.kmk)
    employer_bpjs = bpjs_employer_lookup(payroll_book)

    rows: list[dict[str, object]] = []
    row = empty_row("workspace", "workspace")
    row.update(name="MWS Payroll Agustus 2026 - Simulasi TA 2026/2027", current_period="Agustus 2026", proposed_period="TA 2026/2027")
    rows.append(row)

    matrices = (
        ("current", "regular", "mws-2025-2026", "Matrix MWS 2025/2026", "2025-07-01", current_settings, current_matrix),
        ("proposed", "regular", "mws-2026-2027", "Draft Matrix MWS 2026/2027 Penyesuaian", "2026-07-01", proposed_settings, proposed_matrix),
        ("current", "admin", "mws-current-admin", "Confirmed August admin payroll cells", "2026-08-01", None, current_admin_matrix),
        ("proposed", "admin", "mws-proposed-admin", "KG 3 Admin", "2026-07-01", None, admin_matrix),
    )
    for scenario, matrix_type, matrix_id, name, effective_date, settings, entries in matrices:
        row = empty_row("matrix", matrix_id)
        row.update(matrix_id=matrix_id, matrix_type=matrix_type, name=name, scenario=scenario, effective_date=effective_date,
                   generator_settings=json.dumps(settings, ensure_ascii=True, separators=(",", ":")) if settings else "")
        rows.append(row)

    for scenario, matrix_type, matrix_id, _name, _date, _settings, entries in matrices:
        note = ("Source matrix cached cells, in rupiah" if matrix_type == "regular" or scenario == "proposed"
                else "Confirmed payroll cells only; no authoritative current admin sheet. Evidence in private audit.")
        for code, salary in sorted(entries.items(), key=lambda item: parse_golongan(item[0])):
            group, category, kmk = parse_golongan(code)
            row = empty_row("matrix_entry", f"{scenario}:{matrix_type}:{code}")
            row.update(matrix_id=matrix_id, scenario=scenario, salary_group=group,
                       professional_category=category, kmk_level=kmk, golongan=code,
                       basic_salary=salary, note=note)
            rows.append(row)

    for code, name, direction, _column in COMPONENTS:
        row = empty_row("component_definition", code)
        row.update(code=code, name=name, category="Imported August payroll", direction=direction,
                   calculation_type="manual", default_value="", taxable=direction == "earning",
                   bpjs_kesehatan=False, bpjs_ketenagakerjaan=False,
                   applies_current=True, applies_proposed=True, active=True, rounding=1,
                   notes="Imported cached amount from mws Gaji Agustus 2026(1).xlsx")
        rows.append(row)

    row = empty_row("component_definition", "BPJS_TK_PEMBERI_KERJA")
    row.update(code="BPJS_TK_PEMBERI_KERJA", name="BPJS TK pemberi kerja (aktual Agustus)",
               category="Imported August payroll", direction="employer_contribution",
               calculation_type="manual", default_value="", taxable=False,
               bpjs_kesehatan=False, bpjs_ketenagakerjaan=False,
               applies_current=True, applies_proposed=True, active=True, rounding=1,
               notes="Imported from BPJSTK employer contribution column; not a native Caroll BPJS rule")
    rows.append(row)

    employees: list[dict[str, object]] = []
    assignments: list[dict[str, object]] = []
    audit_employees: list[dict[str, object]] = []
    component_columns = {code: column for code, _name, _direction, column in COMPONENTS}
    first_column = column_index_from_string("C")
    last_column = column_index_from_string("BM")
    column_offset = {
        column: column_index_from_string(column) - first_column
        for column in {"C", "D", "H", "I", "M", "N", "AE", "AN", "BJ", "BL", "BM", *component_columns.values()}
    }
    for excel_row, full_values in enumerate(payroll_rows, start=7):
        values = full_values[first_column - 1:last_column]
        value = lambda column: values[column_offset[column]]
        employee_id = value("C")
        employee_name = value("D")
        current_code = value("M")
        current_basic = value("N")
        if not (employee_id and employee_name and current_code and isinstance(current_basic, (int, float))):
            continue
        employee_id = str(employee_id).strip()
        employee_name = str(employee_name).strip()
        current_code = canonical_golongan(current_code)
        proposed_code = canonical_golongan(calculated.get(normalize_name(employee_name), current_code))
        proposed_code_reason = "Calculated KMK code" if normalize_name(employee_name) in calculated else "No KMK match; preserved current code"
        if proposed_code not in proposed_matrix:
            proposed_code = current_code
            proposed_code_reason = "Calculated code outside source matrix; preserved current code; review required"

        role = str(value("I") or "").strip()
        current_actual = money(current_basic)
        current_type, classification_reason = classify_current(
            policy_key(employee_name), current_code, current_actual, current_matrix, current_admin_matrix, policy)
        proposed_type = current_type
        proposed_reason = "Preserved evidence-based current classification"
        if policy_key(employee_name) == policy["proposed_admin_key"] and proposed_code == "3-PM1":
            proposed_type = "admin"
            proposed_reason = "Explicit user policy for proposed 3-PM1"
        proposed_override: object = ""
        if policy_key(employee_name) == policy["regular_exception_key"]:
            proposed_type = "regular"
            # Preserve the old proposal amount without asserting an admin classification.
            if "admin" in role.lower() and proposed_code in admin_matrix:
                proposed_override = admin_generator[proposed_code]
                proposed_reason = "Regular policy exception; preserved legacy proposal amount pending policy verification"
        current_matrix_salary = {"regular": current_matrix, "admin": current_admin_matrix}.get(current_type, {}).get(current_code)
        current_override: object = "" if current_matrix_salary == current_actual else current_actual
        if policy_key(employee_name) == policy["regular_exception_key"]:
            current_override = current_actual
        override_reason = classification_reason
        if current_override != "":
            override_reason += " Current explicit payroll basic override preserves actual August amount."
        if proposed_override != "":
            override_reason += " " + proposed_reason
        proposed_salary = (proposed_override if proposed_override != "" else
                           {"regular": proposed_matrix, "admin": admin_matrix}.get(proposed_type, {}).get(proposed_code))

        employee = empty_row("employee", employee_id)
        employee.update(
            employee_id=employee_id,
            name=employee_name,
            current_golongan=current_code,
            proposed_golongan=proposed_code,
            current_matrix_type=current_type,
            proposed_matrix_type=proposed_type if proposed_type != current_type else "",
            unit="MWS",
            department=str(value("H") or "").strip(),
            position=role,
            employment_status="",
            join_date="",
            active=True,
            current_basic_override=current_override,
            proposed_basic_override=proposed_override,
            ptkp_status=str(value("AE") or "").strip(),
            bpjs_kesehatan=True,
            bpjs_ketenagakerjaan=True,
            pph_method="gross",
            pph_rate=0,
            pph_fixed_override="",
            notes=override_reason,
        )
        employees.append(employee)

        earning_total = 0
        deduction_total = 0
        for component_code, column in component_columns.items():
            amount = money(value(column))
            if amount == 0:
                continue
            assignment = empty_row("employee_component", f"{employee_id}:{component_code}")
            assignment.update(employee_id=employee_id, component_code=component_code,
                              current_value=amount, proposed_value=amount)
            assignments.append(assignment)
            direction = next(item[2] for item in COMPONENTS if item[0] == component_code)
            if direction == "earning":
                earning_total += amount
            else:
                deduction_total += amount

        employer_amount = employer_bpjs.get(employee_id, 0)
        if employer_amount:
            assignment = empty_row("employee_component", f"{employee_id}:BPJS_TK_PEMBERI_KERJA")
            assignment.update(employee_id=employee_id, component_code="BPJS_TK_PEMBERI_KERJA",
                              current_value=employer_amount, proposed_value=employer_amount)
            assignments.append(assignment)

        source_gross = money(value("AN"))
        source_deductions = money(value("BJ")) + money(value("BL"))
        source_thp = money(value("BM"))
        audit_employees.append({
            "employee_id": employee_id,
            "name": employee_name,
            "current_golongan": current_code,
            "proposed_golongan": proposed_code,
            "source_basic": current_actual,
            "mapped_basic": current_actual if current_override != "" else current_matrix_salary,
            "source_gross": source_gross,
            "mapped_gross": current_actual + earning_total,
            "source_deductions": source_deductions,
            "mapped_deductions": deduction_total,
            "source_thp": source_thp,
            "mapped_thp": current_actual + earning_total - deduction_total,
            "employer_bpjs": employer_amount,
            "current_matrix_type": current_type,
            "proposed_matrix_type": proposed_type,
            "classification_reason": classification_reason,
            "proposed_classification_reason": proposed_reason,
            "proposed_golongan_reason": proposed_code_reason,
            "current_basic_override": current_override,
            "proposed_basic_override": proposed_override,
            "proposed_basic": proposed_salary,
            "proposed_gross": None if proposed_salary is None else proposed_salary + earning_total,
            "proposed_deductions": deduction_total,
            "proposed_thp": None if proposed_salary is None else proposed_salary + earning_total - deduction_total,
            "payroll_row": excel_row,
        })

    rows.extend(employees)
    rows.extend(assignments)

    global_row = empty_row("global_rule", "global_rule")
    global_row.update(currency="IDR", rounding=1, percentage_precision=2,
                      include_inactive=False, allow_negative_thp=False,
                      proposed_defaults_current=True, proposed_matrix_type_defaults_current=True, tax_enabled=False,
                      tax_basis="taxable", tax_rounding=1)
    rows.append(global_row)

    args.output.parent.mkdir(parents=True, exist_ok=True)
    with args.output.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=WORKSPACE_HEADERS, lineterminator="\r\n")
        writer.writeheader()
        writer.writerows(rows)

    mismatches = []
    for employee in audit_employees:
        for metric in ("basic", "gross", "deductions", "thp"):
            if employee[f"source_{metric}"] != employee[f"mapped_{metric}"]:
                mismatches.append({
                    "employee_id": employee["employee_id"],
                    "name": employee["name"],
                    "metric": metric,
                    "source": employee[f"source_{metric}"],
                    "mapped": employee[f"mapped_{metric}"],
                    "difference": employee[f"mapped_{metric}"] - employee[f"source_{metric}"],
                })

    totals = {}
    for metric in ("basic", "gross", "deductions", "thp"):
        totals[metric] = {
            "source": sum(employee[f"source_{metric}"] for employee in audit_employees),
            "mapped": sum(employee[f"mapped_{metric}"] for employee in audit_employees),
        }
        totals[metric]["difference"] = totals[metric]["mapped"] - totals[metric]["source"]

    report = {
        "workspace_csv": str(args.output),
        "employee_count": len(employees),
        "current_matrix_entries": len(current_matrix),
        "proposed_matrix_entries": len(proposed_matrix),
        "matrix_generator_settings": {"current": current_settings, "proposed": proposed_settings},
        "matrix_generator_cell_mismatches": matrix_mismatches,
        "calculated_golongan_matches": sum(1 for employee in audit_employees if employee["current_golongan"] != employee["proposed_golongan"]),
        "matrix_counts": {f"{s}:{t}": len(entries) for s, t, _id, _name, _date, _settings, entries in matrices},
        "type_counts": {scenario: dict(Counter(e[f"{scenario}_matrix_type"] or "unresolved" for e in audit_employees)) for scenario in ("current", "proposed")},
        "classifications": audit_employees,
        "unresolved_classifications": [e for e in audit_employees if not e["current_matrix_type"] or not e["proposed_matrix_type"] or e["proposed_basic"] is None],
        "overrides": {scenario: sum(e[f"{scenario}_basic_override"] != "" for e in audit_employees) for scenario in ("current", "proposed")},
        "current_admin_evidence": admin_evidence,
        "proposed_admin_generator_cell_mismatches": [code for code, salary in admin_matrix.items() if admin_generator.get(code) != salary],
        "source_paths": {key: str(getattr(args, key).resolve()) for key in ("payroll", "kmk", "matrix")},
        "cell_reconciliation": [
            {"scenario": scenario, "matrix_type": matrix_type, "golongan": code,
             "mapped": salary, "source": salary,
             "evidence": admin_evidence[code] if scenario == "current" and matrix_type == "admin" else {
                 "sheet": ("KG 3 Admin" if matrix_type == "admin" else
                           "Matrix MWS 25_26" if scenario == "current" else "Draft Matrix MWS 26_27 Penyesua"),
                 "row": (12 if scenario == "current" else 13) + parse_golongan(code)[2] - 1,
                 "column": (3 + (int(parse_golongan(code)[0]) - 1) * 6
                            + ("PM", "P", "M", "U").index(parse_golongan(code)[1])),
                 "unit_multiplier": 1000,
             }}
            for scenario, matrix_type, _id, _name, _date, _settings, entries in matrices
            for code, salary in entries.items()
        ],
        "proposed_resolved_subtotals": {
            "employee_count": sum(e["proposed_basic"] is not None for e in audit_employees),
            **{metric: sum(e[f"proposed_{metric}"] for e in audit_employees if e["proposed_basic"] is not None)
               for metric in ("basic", "gross", "deductions", "thp")},
        },
        "scenario_totals": {
            "current": {metric: totals[metric]["mapped"] for metric in totals},
            "proposed": {metric: (sum(e[f"proposed_{metric}"] for e in audit_employees) if all(e[f"proposed_{metric}"] is not None for e in audit_employees) else None) for metric in totals},
        },
        "component_assignments": len(assignments),
        "totals": totals,
        "mismatches": mismatches,
        "limitations": [
            "PPh 21 is disabled because the August payroll sheet has no employee PPh column to import.",
            "Actual employee and employer BPJS are imported as manual payroll components because Caroll has no per-employee BPJS amount or wage-basis override.",
            "Current admin has no authoritative sheet: only user-confirmed payroll cells are included, with source evidence; other cells remain absent.",
            "Regular classification accepts exact salary matches or nearest-thousand source salary matches; explicit overrides retain every actual current basic.",
            "Unmatched or ambiguous classifications remain blank and block core payroll; position text alone never establishes a type.",
            "Configured source cache differences are retained, not patched into payroll components.",
            "Proposed non-basic components and deductions are held equal to August values; only matrix/Golongan-driven basic salary changes in the projection.",
        ],
    }
    report["reconciliation_passed"] = reconciliation_passed(totals, mismatches, policy["reconciliation"])
    for employee in report["unresolved_classifications"]:
        code = employee["current_golongan"]
        proposed_code = employee["proposed_golongan"]
        candidates = {
            "current_regular": source_current_matrix.get(code),
            "current_admin_confirmed": current_admin_matrix.get(code),
            "proposed_regular": source_proposed_matrix.get(proposed_code),
            "proposed_admin": admin_matrix.get(proposed_code),
        }
        employee["candidate_evidence"] = {
            label: {"salary": salary,
                    "nearest_thousand": None if salary is None else money(Decimal(salary) / 1000) * 1000,
                    "matches_actual_exactly": salary == employee["source_basic"],
                    "matches_actual_rounded": salary is not None and money(Decimal(salary) / 1000) * 1000 == employee["source_basic"]}
            for label, salary in candidates.items()
        }
    report["unresolved_count"] = len(report["unresolved_classifications"])
    report["results_exported"] = False
    report["policy_config"] = str(args.policy_config.resolve())
    report["results_csv"] = str(results_path)
    report["audit_json"] = str(args.report)
    args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    if not report["reconciliation_passed"]:
        raise ValueError("August reconciliation failed; inspect private audit")
    # Use the same normalization/validation/calculation APIs as the UI. Capture
    # all output: core issues may contain identities and belong only in the audit.
    with tempfile.TemporaryDirectory(prefix=".august-results-", dir=args.output.parent) as staging:
        staged_result = Path(staging) / results_path.name
        return finish_build(args, report, results_path, staged_result)


def finish_build(args, report, results_path, staged_result):
    integration = subprocess.run([
        "node", "-e", """
        const fs = require('node:fs');
        const C = require('./js/state.js');
        require('./js/payroll.js'); require('./js/csv.js');
        try {
          const ws = C.csv.importWorkspace(fs.readFileSync(process.argv[1], 'utf8'));
          const result = C.calculatePayroll(ws);
          if (result.totals) fs.writeFileSync(process.argv[2], C.csv.exportResults(result));
          process.stdout.write(JSON.stringify({totals: result.totals, issues: result.issues}));
        } catch (error) {
          process.stdout.write(JSON.stringify({error: error.message}));
          process.exitCode = 1;
        }
        """, str(args.output.resolve()), str(staged_result.resolve())
    ], cwd=Path(__file__).resolve().parents[1], capture_output=True, text=True, timeout=60)
    report["core"] = json.loads(integration.stdout) if integration.stdout else {"error": "Core integration failed"}
    report["blocking_error_count"] = sum(issue.get("severity") == "error" for issue in report["core"].get("issues", []))
    success = (integration.returncode == 0 and bool(report["core"].get("totals"))
               and not report["core"].get("error") and not report["blocking_error_count"]
               and staged_result.is_file())
    report["results_exported"] = False
    args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    if not success:
        raise ValueError("Core payroll blocked; inspect private audit")
    try:
        os.replace(staged_result, results_path)
        report["results_exported"] = True
        args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    except Exception:
        results_path.unlink(missing_ok=True)
        raise
    return report


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--payroll", type=Path, required=True)
    parser.add_argument("--kmk", type=Path, required=True)
    parser.add_argument("--matrix", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--report", type=Path, required=True)
    parser.add_argument("--policy-config", type=Path, required=True)
    args = parser.parse_args()
    try:
        report = build(args)
    except Exception:
        # Even parsing errors can contain private cell values. Never echo them.
        parser.exit(1, "Workspace build failed; inspect inputs and private audit if available.\n")
    print(json.dumps({key: report[key] for key in ("reconciliation_passed", "results_exported")}))


if __name__ == "__main__":
    main()
